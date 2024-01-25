from openpyxl.styles import Font, Alignment
from selenium import webdriver
from selenium.webdriver.common.by import By
import openpyxl 
from openpyxl.chart import PieChart, Reference 
#entrar no site
driver = webdriver.Chrome()
driver.get('https://www.kabum.com.br/promocao/MENU_PCGAMER')
#extrair titulo
titulos = driver.find_elements(By.XPATH, " //span[@class='sc-d79c9c3f-0 nlmfp sc-cdc9b13f-16 eHyEuD nameCard']")
#extrair preco
precos = driver.find_elements(By.XPATH, "//span[@class='sc-620f2d27-2 bMHwXA priceCard']")
#criar plhanilha no excel
workbook = openpyxl.Workbook()
#criando página produtos
workbook.create_sheet('produtos')
#selecionando a página produtos
sheet_produtos = workbook['produtos']
sheet_produtos['A1'].value = 'Produto'
sheet_produtos['B1'].value = 'Preço'
#inserir dados na planilha
for titulo, preco in zip(titulos, precos):
    
    preco_numerico = float(preco.text.replace('R$', '').replace(',', '.'))
    sheet_produtos.append([titulo.text, preco_numerico])


#coletando os preços e fazendo uma média

sheet_produtos['B23']='=AVERAGE(B2:B21)'
#gerando um gráfico
chart = PieChart() 
#lendo os limites de linhas e colunas
labels = Reference(sheet_produtos, min_col = 1, 
                   min_row = 2, max_row = 21) 
#lendo os dados das linhas e colunas  
data = Reference(sheet_produtos, min_col = 2, 
                   min_row = 1, max_row = 21) 
chart.add_data(data, titles_from_data = True) 
chart.set_categories(labels) 
chart.title = " PIE-CHART "
#Selecionando em que coluna o gráfico irá iniciar
sheet_produtos.add_chart(chart, "N2")
#salvar planilha
workbook.save('planilha_automatizada.xlsx')




